import csv
import io
import json
import os
import re
import secrets
import smtplib
import sqlite3
import uuid
from collections import Counter
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from email.utils import formataddr
from functools import wraps
from urllib.parse import urlencode

from flask import (Flask, Response, abort, flash, g, redirect, render_template, request, send_from_directory, session,
                   url_for)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps, UnidentifiedImageError
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "alumni.db")
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
ALLOWED_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic"}
# Event pictures are public (shown on the landing page), so they live under static/
EVENT_IMAGE_DIR = os.path.join(BASE_DIR, "static", "event_images")
EVENT_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
EVENT_IMAGE_MAX_SIZE = 1600   # pixels; larger pictures are shrunk so the landing page loads quickly

# All pages (index.html, admin.html, ...) live in the templates/ folder
app = Flask(__name__, template_folder=os.path.join(BASE_DIR, "templates"))
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10 MB limit for the proof of employment upload


def load_secret_key():
    """Keep the admin login working across restarts by storing the session key in a file."""
    key_path = os.path.join(BASE_DIR, ".secret_key")
    if not os.path.exists(key_path):
        with open(key_path, "w") as f:
            f.write(secrets.token_hex(32))
    with open(key_path) as f:
        return f.read().strip()


app.secret_key = load_secret_key()
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"  # don't send the login cookie with forms posted from other sites

# Email account used to send password-change verification codes (set these before starting the app).
# For Gmail: MAIL_SERVER=smtp.gmail.com, MAIL_PORT=587, MAIL_USERNAME=<address>, MAIL_PASSWORD=<16-letter app password>
# If MAIL_SERVER is empty, codes are printed in the terminal running Flask instead of being emailed.
# The settings can also be typed into mail_settings.txt (next to this file), one NAME=value per line.
def load_mail_settings():
    path = os.path.join(BASE_DIR, "mail_settings.txt")
    if not os.path.exists(path):
        return
    settings = {}
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            name, separator, value = line.strip().partition("=")
            if separator and name.strip().startswith("MAIL_") and value.strip():
                settings[name.strip()] = value.strip()
    # Until the account and its password are filled in, the file is ignored (codes go to the terminal)
    if settings.get("MAIL_USERNAME") and settings.get("MAIL_PASSWORD"):
        for name, value in settings.items():
            os.environ.setdefault(name, value)   # a real environment variable wins


load_mail_settings()
MAIL_SERVER = os.environ.get("MAIL_SERVER", "")
MAIL_PORT = int(os.environ.get("MAIL_PORT", "587"))
MAIL_USERNAME = os.environ.get("MAIL_USERNAME", "")
MAIL_PASSWORD = os.environ.get("MAIL_PASSWORD", "")
MAIL_FROM = os.environ.get("MAIL_FROM", MAIL_USERNAME)

MIN_PASSWORD_LENGTH = 8
CODE_LIFETIME = timedelta(minutes=10)   # how long a verification code stays valid
MAX_CODE_ATTEMPTS = 5                   # wrong codes allowed before the change is cancelled

# The colleges that take part in the tracer study: (short code, full name as saved by the survey).
# Each coordinator belongs to one college and only sees that college's responses, graduates, emails and events.
COLLEGES = sorted([
    ("CCI", "College of Computing and Informatics"),
    ("CAS", "College of Arts and Sciences"),
    ("COED", "College of Education"),
    ("CIT", "College of Industrial Technology"),
    ("CGBE", "College of Global Business and Enterprise"),
    ("CEA", "College of Engineering and Architecture"),
], key=lambda college: college[1])   # shown everywhere in alphabetical order
COLLEGE_NAMES = dict(COLLEGES)
# The college that the site started with: older records without a college belong to it
DEFAULT_COLLEGE = "CCI"

# Degree programs of the College of Computing and Informatics: (short code, full name as saved by the survey).
# These are the starting list; every college's programs are kept in the "programs" table (Programs page).
PROGRAMS = [
    ("BSIT", "BACHELOR OF SCIENCE IN INFORMATION TECHNOLOGY"),
    ("BSIS", "BACHELOR OF SCIENCE IN INFORMATION SYSTEMS"),
    ("BSCS", "BACHELOR OF SCIENCE IN COMPUTER SCIENCE"),
    ("MSCS-SE", "MASTER OF SCIENCE IN COMPUTER SCIENCE SOFTWARE ENGINEERING TRACK"),
    ("MSCS-NET", "MASTER OF SCIENCE IN COMPUTER SCIENCE NETWORKING TRACK"),
]

# Short description of each program for the landing page: college code -> program code -> (what it is about,
# where graduates work). A program added later on the Programs page shows just its name until it is described here.
PROGRAM_INFO = {
    "CCI": {
        "BSIT": ("Applies computing technology to the everyday needs of organizations: building and managing networks, "
                 "systems, websites and mobile apps, and keeping IT services secure and running.",
                 "Web and mobile developer, network or systems administrator, IT support and cybersecurity specialist"),
        "BSIS": ("Bridges technology and business: designing and managing information systems that support an "
                 "organization's operations, data and decision-making.",
                 "Systems or business analyst, database administrator, IT project coordinator, ERP specialist"),
        "BSCS": ("Studies the science behind computing: algorithms, programming, data and intelligent systems, "
                 "to create new software and solve complex computing problems.",
                 "Software engineer, data scientist, AI and machine learning developer, researcher"),
        "MSCS-SE": ("Advanced study of building large, reliable software: requirements, architecture, testing and "
                    "managing software projects, plus research in software engineering.",
                    "Senior software engineer, software architect, project lead, educator and researcher"),
        "MSCS-NET": ("Advanced study of computer networks: designing, securing and managing networks and cloud "
                     "infrastructure, plus research in networking and communications.",
                     "Network architect or engineer, security specialist, cloud engineer, educator and researcher"),
    },
    "CAS": {
        "MA-MATH": ("Advanced study of pure and applied mathematics: analysis, algebra and mathematical modelling, "
                    "with research that prepares graduates to teach and use mathematics at a higher level.",
                    "College mathematics instructor, researcher, statistician, data or quantitative analyst"),
        "BSEL": ("Studies how the English language works and how it is used: linguistics, professional and technical "
                 "writing, and communication across media and cultures.",
                 "Writer or editor, communications officer, language trainer, content specialist, researcher"),
        "BHS": ("Prepares graduates to help individuals, families and communities: counselling basics, case "
                "management, and planning and running social and community service programs.",
                "Human services or case worker, community program officer, guidance staff, welfare assistant"),
        "BSBIO-BIOTECH": ("Studies living systems with a focus on biotechnology: genetics, microbiology and molecular "
                          "techniques used in health, agriculture, industry and the environment.",
                          "Laboratory analyst, research assistant, quality control officer, biotechnologist; "
                          "preparation for medicine"),
        "BSCD": ("Trains graduates to work with communities: organizing people, assessing needs, and planning, "
                 "managing and evaluating development projects.",
                 "Community development officer, project coordinator, NGO or local government program staff"),
        "BSMATH": ("Builds strong skills in mathematical reasoning, statistics and computing to model and solve "
                   "problems in science, business and industry.",
                   "Data or statistical analyst, actuarial assistant, researcher, mathematics instructor"),
    },
    "COED": {
        "EDD": ("The highest degree in education: advanced leadership, policy and research for those who lead "
                "schools, programs and educational change.",
                "School head or superintendent, dean, education program supervisor, professor and researcher"),
        "MAED-EM": ("Graduate study of how schools are led and run: supervision, planning, school finance and "
                    "personnel, backed by educational research.",
                    "Principal or school head, department chair, education supervisor, college instructor"),
        "MTLED-AFA": ("Graduate study for teachers of technology and livelihood education, with advanced content, "
                      "teaching methods and research in agriculture and fishery arts.",
                      "Master teacher, TLE department head, technical-vocational trainer, college instructor"),
        "MTLED-HE": ("Graduate study for teachers of technology and livelihood education, with advanced content, "
                     "teaching methods and research in home economics.",
                     "Master teacher, TLE department head, technical-vocational trainer, college instructor"),
        "MTLED-IA": ("Graduate study for teachers of technology and livelihood education, with advanced content, "
                     "teaching methods and research in industrial arts.",
                     "Master teacher, TLE department head, technical-vocational trainer, college instructor"),
        "MTLED-ICT": ("Graduate study for teachers of technology and livelihood education, with advanced content, "
                      "teaching methods and research in information and communication technology.",
                      "Master teacher, ICT coordinator, TLE department head, technical-vocational trainer"),
        "BEED": ("Prepares teachers for the elementary grades: child development, teaching methods across all "
                 "subjects, assessment and classroom practice.",
                 "Elementary school teacher, tutor, learning facilitator, curriculum or training assistant"),
        "BSED-FIL": ("Prepares high school teachers of Filipino: the language, its literature, and the methods and "
                     "assessment used to teach them.",
                     "Secondary school Filipino teacher, writer or editor, translator, language tutor"),
        "BSED-MATH": ("Prepares high school teachers of mathematics: algebra, geometry, statistics and calculus, "
                      "with the methods and assessment used to teach them.",
                      "Secondary school mathematics teacher, tutor or reviewer, training or testing staff"),
        "BSED-SCI": ("Prepares high school teachers of science: biology, chemistry, physics and earth science, "
                     "with laboratory work and the methods used to teach them.",
                     "Secondary school science teacher, laboratory instructor, science tutor, training staff"),
        "BTVTED-AT": ("Prepares teachers and trainers of automotive technology for technical-vocational schools: "
                      "vehicle servicing skills together with training methods and assessment.",
                      "Technical-vocational teacher, TESDA trainer or assessor, automotive service technician"),
        "BTVTED-BCW": ("Prepares teachers and trainers of beauty care and wellness for technical-vocational schools: "
                       "hair, skin and wellness services together with training methods and assessment.",
                       "Technical-vocational teacher, TESDA trainer or assessor, salon or spa practitioner"),
        "BTVTED-CP": ("Prepares teachers and trainers of computer programming for technical-vocational schools: "
                      "software development skills together with training methods and assessment.",
                      "Technical-vocational or ICT teacher, TESDA trainer or assessor, junior programmer"),
        "BTVTED-DT": ("Prepares teachers and trainers of drafting technology for technical-vocational schools: "
                      "technical drawing and computer-aided design together with training methods and assessment.",
                      "Technical-vocational teacher, TESDA trainer or assessor, CAD operator or draftsman"),
        "BTVTED-ELT": ("Prepares teachers and trainers of electrical technology for technical-vocational schools: "
                       "electrical installation and maintenance together with training methods and assessment.",
                       "Technical-vocational teacher, TESDA trainer or assessor, electrical technician"),
        "BTVTED-ELX": ("Prepares teachers and trainers of electronics technology for technical-vocational schools: "
                       "electronic circuits, servicing and assembly together with training methods and assessment.",
                       "Technical-vocational teacher, TESDA trainer or assessor, electronics technician"),
        "BTVTED-FSM": ("Prepares teachers and trainers of food and service management for technical-vocational "
                       "schools: cookery and food service operations together with training methods and assessment.",
                       "Technical-vocational teacher, TESDA trainer or assessor, food service supervisor"),
        "BTLED-HE": ("Prepares teachers of Technology and Livelihood Education for junior high school, with home "
                     "economics as the area of focus: foods, clothing, and home and family living.",
                     "TLE teacher, livelihood skills trainer, home economics practitioner or entrepreneur"),
        "BTLED-IA": ("Prepares teachers of Technology and Livelihood Education for junior high school, with "
                     "industrial arts as the area of focus: drafting, woodwork, metalwork, electricity and electronics.",
                     "TLE teacher, livelihood skills trainer, shop or industrial arts instructor"),
        "BTLED-ICT": ("Prepares teachers of Technology and Livelihood Education for junior high school, with "
                      "information and communication technology as the area of focus.",
                      "TLE or ICT teacher, school ICT coordinator, computer skills trainer"),
        "CT": ("A short program of professional education courses for graduates of other degrees who want to "
               "teach and qualify for the Licensure Examination for Teachers.",
               "Licensed teacher in basic education, tutor, trainer in the graduate's own field"),
    },
    "CIT": {
        "DIT": ("The highest degree in industrial technology: advanced technical leadership, innovation and applied "
                "research for industry and technical education.",
                "Industry executive or consultant, technical education administrator, professor and researcher"),
        "MIT": ("Graduate study in industrial technology: managing production and technical operations, advanced "
                "practice in a technology area, and applied research.",
                "Production or plant supervisor, technical manager, trainer, college instructor"),
        "BIT-AD": ("Hands-on training in architectural drafting: preparing building plans and working drawings by "
                   "hand and with computer-aided design, plus supervision skills for industry.",
                   "Architectural draftsman, CAD operator, estimator, site documentation staff"),
        "BIT-AT": ("Hands-on training in automotive technology: diagnosing, servicing and repairing engines, "
                   "electrical systems and chassis, plus supervision skills for industry.",
                   "Automotive technician, service advisor, shop supervisor, parts and service specialist"),
        "BIT-CT": ("Hands-on training in construction technology: building methods, materials, estimating and site "
                   "practice, plus supervision skills for industry.",
                   "Construction foreman, site supervisor, estimator, building maintenance technician"),
        "BIT-ELT": ("Hands-on training in electrical technology: wiring, motors and controls, and power "
                    "installation and maintenance, plus supervision skills for industry.",
                    "Electrical technician, industrial electrician, maintenance supervisor, installation estimator"),
        "BIT-ELX": ("Hands-on training in electronics technology: circuits, instrumentation, communication and "
                    "control systems, plus supervision skills for industry.",
                    "Electronics technician, instrumentation or test technician, production line supervisor"),
        "BIT-HVACR": ("Hands-on training in heating, ventilating, air conditioning and refrigeration: installing, "
                      "servicing and troubleshooting cooling systems, plus supervision skills for industry.",
                      "HVAC/R technician, building or plant maintenance staff, service supervisor"),
        "BIT-MT": ("Hands-on training in mechanical technology: machining, fabrication and maintenance of "
                   "machines and equipment, plus supervision skills for industry.",
                   "Machinist, mechanical maintenance technician, production or quality supervisor"),
        "BIT-BCW": ("Hands-on training in beauty care and wellness: hair, skin and nail care and wellness "
                    "services, plus the skills to manage a salon or spa.",
                    "Cosmetologist, salon or spa supervisor, wellness practitioner, business owner"),
        "BIT-CUL": ("Hands-on training in culinary technology: food preparation, baking and kitchen operations, "
                    "with food safety and the skills to supervise a kitchen.",
                    "Cook or chef, kitchen supervisor, baker, food business owner"),
        "BIT-WF": ("Hands-on training in welding and fabrication: welding processes, metalwork and reading "
                   "fabrication drawings, plus supervision skills for industry.",
                   "Welder or fabricator, welding inspector, fabrication shop supervisor"),
        "BIT-AF": ("Hands-on training in apparel and fashion technology: pattern making, garment construction and "
                   "production methods, plus supervision skills for the garments industry.",
                   "Pattern maker, garment production supervisor, quality controller, dressmaker or tailor"),
        "BSFDM-MW": ("Combines design and business for the fashion industry, with a focus on men's wear: design, "
                     "pattern making, tailoring, and merchandising and retail.",
                     "Fashion designer, tailor, merchandiser or buyer, stylist, clothing business owner"),
        "BSFDM-WW": ("Combines design and business for the fashion industry, with a focus on women's wear: design, "
                     "pattern making, dressmaking, and merchandising and retail.",
                     "Fashion designer, dressmaker, merchandiser or buyer, stylist, clothing business owner"),
    },
    "CGBE": {
        "BSENTREP": ("Teaches how to start and grow a business: spotting opportunities, planning, marketing, "
                     "finance and managing a small enterprise.",
                     "Business owner, business development officer, marketing or sales staff, enterprise consultant"),
        "BSHM-CA": ("Prepares graduates for the hotel and restaurant industry, with a focus on culinary arts: "
                    "professional cooking, kitchen management and food service operations.",
                    "Chef or cook, kitchen supervisor, restaurant or catering manager, food business owner"),
        "BSHM-CSS": ("Prepares graduates for the hotel and restaurant industry, with a focus on service aboard "
                     "cruise ships: housekeeping, food and beverage and guest services at sea.",
                     "Cruise ship steward or guest service staff, hotel or restaurant supervisor, front office staff"),
        "BSTM": ("Studies the travel and tourism industry: planning tours and destinations, travel services, "
                 "events and sustainable tourism.",
                 "Travel or tour officer, tourism officer, events coordinator, airline or hotel guest service staff"),
    },
    "CEA": {
        "BSARCH": ("Trains graduates to design buildings and spaces that are safe, functional and beautiful: "
                   "architectural design, building technology, and planning.",
                   "Architect (after licensure), architectural designer, project or site architect, planner"),
        "BSCE": ("Covers the design and construction of roads, bridges, buildings and water systems: structural, "
                 "geotechnical, transportation and water resources engineering.",
                 "Civil engineer (after licensure), structural designer, site or project engineer, estimator"),
        "BSEE": ("Covers the generation, distribution and use of electric power: power systems, machines, "
                 "controls and electrical design.",
                 "Electrical engineer (after licensure), power plant or distribution engineer, design engineer"),
        "BSME": ("Covers the design, manufacture and maintenance of machines and energy systems: thermodynamics, "
                 "machine design, and power and industrial plants.",
                 "Mechanical engineer (after licensure), plant or maintenance engineer, design engineer"),
        "BSECE": ("Covers electronic devices and systems: circuits, communications, signal processing and "
                  "embedded and control systems.",
                  "Electronics engineer (after licensure), telecommunications or network engineer, systems designer"),
    },
}


@app.template_filter("college_logo")
def college_logo(code):
    """Address of the college's logo, or '' if it has none yet.

    To add a logo, save it as static/images/colleges/<college code in small letters>.png, e.g. cgbe.png
    """
    filename = f"{(code or '').lower()}.png"
    if code in COLLEGE_NAMES and os.path.exists(os.path.join(BASE_DIR, "static", "images", "colleges", filename)):
        return f"/static/images/colleges/{filename}"
    return ""


@app.template_filter("college_name")
def college_name(code):
    """'CAS' -> 'College of Arts and Sciences' (used to label which college posted an event)."""
    return COLLEGE_NAMES.get(code, code or "")


@app.template_filter("program_name")
def program_name(name):
    """'BACHELOR OF SCIENCE IN ...' -> 'Bachelor of Science in ...' for display. Codes like 'MSCS' stay as they are."""
    if " " not in name.strip():
        return name
    small_words = {"of", "in", "and", "the", "with"}
    words = name.lower().split()
    # "(biotechnology)" -> "(Biotechnology)": capitalize the first letter, not the bracket
    capitalize = lambda w: w[0] + w[1:].capitalize() if w[:1] == "(" else w.capitalize()
    return " ".join(w if w in small_words and i else capitalize(w) for i, w in enumerate(words))

# Form field name (HTML name="...") -> column in the tracer table.
# When you add a question to index.html, add its name here too.
FIELDS = [
    # Step 1: Data Privacy Consent and Demographics
    "respondentType",   # "Alumni" (index.html) or "New Graduate" (survey_new_graduate.html), from a hidden field
    "consent", "lname", "fname", "mname", "sex", "mnum", "email", "address",
    "bdate", "civilStatus",
    # Step 2: Educational Background and Professional Examinations
    "college", "section", "program", "yeargrad", "yeargradOther", "honor", "reason",
    "examName", "examDate", "examRating",
    # Step 3: Trainings and Advance Studies
    "trainingTitle", "trainingDuration", "trainingInstitution", "trainingStatus",
    "advReason", "postbaccReason",
    # Step 4: Employment Data
    "employed", "supervisorEmail", "unempReason", "empStatus", "occupation",
    "jobLevel", "currentSalary", "jobAligned", "higherEdRequired", "govEmployed",
    "company", "businessLine", "workPlace", "firstJob", "stayReason",
    "firstJobRelated", "acceptReason", "firstJobStay", "findFirstJob", "timeToJob",
    "firstJobLevel", "firstJobSalary", "leftReason",
    # Step 5: Curricular Improvement and Entrepreneurial Venture
    "curriculumRelevant", "competencies", "selfEmpCompetencies", "suggestions",
    "businessNature", "sideHustle", "sideHustleDetails", "startup",
    # Step 6: Recommendation and Confirmation
    "recommendIsatu", "recommendCourse", "proudOf", "otherRecommendations",
    "confirmation",
    # Only in the New Graduate registration form
    "birthPlace", "citizenship", "permanentAddress", "facebook", "educLevel",
]

# Columns filled in by the server rather than typed into the form
EXTRA_COLUMNS = ["proofFile", "submittedAt"]

# How responses are shown on the All Responses page and in downloads: (section, [(column, label), ...])
RECORD_GROUPS = [
    ("Record", [
        ("submittedAt", "Date Submitted"), ("respondentType", "Respondent Type"), ("consent", "Data Privacy Consent"),
]),
    ("Demographics", [
        ("lname", "Last Name"), ("fname", "First Name"), ("mname", "Middle Name"),
        ("sex", "Sex"), ("mnum", "Contact Number"), ("email", "Email"),
        ("address", "Present Address"), ("bdate", "Date of Birth"), ("civilStatus", "Civil Status"),
        ("birthPlace", "Place of Birth"), ("citizenship", "Citizenship"), ("permanentAddress", "Permanent Address"),
        ("facebook", "Facebook Account"),
]),
    ("Educational Background", [
        ("educLevel", "Current Educational Level"), ("college", "College"), ("section", "Section"),
        ("program", "Course and Specialization"),
        ("yeargrad", "Year Graduated"), ("yeargradOther", "Year Graduated (Others)"),
        ("honor", "Honors and Awards"), ("reason", "Reasons for Taking the Course"),
]),
    ("Professional Examinations Passed", [
        ("examName", "Name of Examination"), ("examDate", "Date Taken"), ("examRating", "Rating (%)"),
]),
    ("Trainings and Advance Studies", [
        ("trainingTitle", "Title of Training / Advance Study"), ("trainingDuration", "Duration and Credits Earned"),
        ("trainingInstitution", "Institution"), ("trainingStatus", "Status"),
        ("advReason", "Reasons for Pursuing Advance Studies"),
        ("postbaccReason", "Reasons for Taking the Post-Baccalaureate Program"),
]),
    ("Employment Data", [
        ("employed", "Currently Employed"), ("supervisorEmail", "Supervisor's Company Email"),
        ("unempReason", "Reasons Not Yet Employed"), ("empStatus", "Employment Status"),
        ("occupation", "Present Job Position"), ("jobLevel", "Present Job Level"),
        ("currentSalary", "Current Gross Monthly Earning"),
        ("jobAligned", "Employment Aligned to Course"), ("higherEdRequired", "Job Requires Higher Education"),
        ("govEmployed", "Government or Private"), ("company", "Company / Institution"),
        ("businessLine", "Major Line of Business"), ("workPlace", "Work Location"),
        ("firstJob", "First Job After College"), ("stayReason", "Reasons for Staying"),
        ("firstJobRelated", "First Job Related to Course"), ("acceptReason", "Reasons for Accepting First Job"),
        ("firstJobStay", "Length of Stay in First Job"), ("findFirstJob", "How First Job Was Found"),
        ("timeToJob", "Time to Land First Job"), ("firstJobLevel", "First Job Level"),
        ("firstJobSalary", "Initial Gross Monthly Earning"), ("leftReason", "Reasons for Leaving First Job"),
]),
    ("Curricular Improvement", [
        ("curriculumRelevant", "Curriculum Relevant to First Job"),
        ("competencies", "Useful Competencies"), ("selfEmpCompetencies", "Useful Competencies (Self-employed)"),
        ("suggestions", "Suggestions to Improve the Curriculum"),
]),
    ("Entrepreneurial Venture", [
        ("businessNature", "Nature of Business (Self-employed)"), ("sideHustle", "Has Side Hustle"),
        ("sideHustleDetails", "Side Hustle / Business Venture"), ("startup", "Created a DTI-registered Startup"),
]),
    ("Recommendation", [
        ("recommendIsatu", "Will Recommend ISAT U"), ("recommendCourse", "Will Recommend Course"),
        ("proudOf", "I Am Proud of ISAT U Because..."), ("otherRecommendations", "Other Recommendations"),
]),
    ("Proof and Confirmation", [
        ("proofFile", "Proof of Employment"), ("confirmation", "Confirmation"),
]),
]

# The two survey forms ask different questions, so each respondent type gets its own set of columns.
NEW_GRADUATE_ONLY = {"birthPlace", "citizenship", "permanentAddress", "facebook", "educLevel"}
RESPONDENT_TYPES = {
    # key in the URL -> (tab label, columns grouped by section)
    "alumni": ("Alumni", [
        (section, [(key, label) for key, label in columns if key not in NEW_GRADUATE_ONLY and key != "respondentType"])
        for section, columns in RECORD_GROUPS
    ]),
    "new-graduate": ("New Graduates", [
        ("Record", [("submittedAt", "Date Submitted"), ("consent", "Data Privacy Consent")]),
        ("Personal Information", [
            ("lname", "Last Name"), ("fname", "First Name"), ("mname", "Middle Initial"), ("email", "Email"),
            ("bdate", "Birthdate"), ("sex", "Sex"), ("civilStatus", "Civil Status"),
            ("birthPlace", "Place of Birth"), ("citizenship", "Citizenship"),
        ]),
        ("Contact Information", [
            ("permanentAddress", "Permanent Address"), ("mnum", "Contact Number"), ("facebook", "Facebook Account"),
        ]),
        ("Education", [
            ("educLevel", "Current Educational Level"), ("yeargrad", "Year Graduated"), ("college", "College"),
            ("program", "Course"),
        ]),
    ]),
}

# Download formats for the responses: key -> (label in the menu, MIME type). The first one is the default.
EXPORT_FORMATS = {
    "xlsx": ("Excel workbook (.xlsx)", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "csv": ("CSV, comma-separated (.csv)", "text/csv; charset=utf-8"),
    "json": ("JSON (.json)", "application/json"),
}


def get_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    """Create the tables if missing and add any tracer columns listed in FIELDS."""
    with get_connection() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS admins (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                full_name TEXT NOT NULL,
                email TEXT NOT NULL UNIQUE,
                mobile TEXT,
                password_hash TEXT NOT NULL,
                created_at TEXT
            )
            """
        )
        # Two kinds of account share this table: the "admin" (system administrator, registers and manages the
        # coordinators) and the "coordinator" (college coordinator, works with the tracer study data).
        columns = {row["name"] for row in conn.execute("PRAGMA table_info(admins)")}
        for column, definition in (("role", "TEXT NOT NULL DEFAULT 'coordinator'"),
                                   ("active", "INTEGER NOT NULL DEFAULT 1"),
                                   ("college", "TEXT")):   # a coordinator's college code; empty for the administrator
            if column not in columns:
                try:
                    conn.execute(f"ALTER TABLE admins ADD COLUMN {column} {definition}")
                except sqlite3.OperationalError as e:
                    if "duplicate column" not in str(e):
                        raise
        # Accounts made before roles existed: the first one becomes the system administrator
        if not conn.execute("SELECT 1 FROM admins WHERE role = 'admin'").fetchone():
            conn.execute("UPDATE admins SET role = 'admin', active = 1 WHERE id = (SELECT MIN(id) FROM admins)")
        # One pending password change per admin, waiting for the emailed code
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS password_change_requests (
                admin_id INTEGER PRIMARY KEY REFERENCES admins(id),
                code_hash TEXT NOT NULL,
                new_password_hash TEXT NOT NULL,
                expires_at TEXT NOT NULL,
                attempts INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        conn.execute("UPDATE admins SET college = ? WHERE role = 'coordinator' AND college IS NULL", (DEFAULT_COLLEGE,))
        # Degree programs of each college, kept by that college's coordinator
        first_run = not conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'programs'").fetchone()
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS programs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                college TEXT NOT NULL,
                code TEXT NOT NULL,
                name TEXT NOT NULL,
                UNIQUE (college, code)
            )
            """
        )
        if first_run:
            conn.executemany("INSERT OR IGNORE INTO programs (college, code, name) VALUES (?, ?, ?)",
                             [(DEFAULT_COLLEGE, code, name) for code, name in PROGRAMS])
        # Total number of graduates per college and year, entered by the coordinator, for the response rate
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS graduate_totals (
                college TEXT NOT NULL,
                year TEXT NOT NULL,
                total INTEGER NOT NULL CHECK (total > 0),
                updated_at TEXT,
                updated_by TEXT,
                PRIMARY KEY (college, year)
            )
            """
        )
        # Totals entered before the other colleges were added (old graduate_counts table) belong to CCI
        if conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'graduate_counts'").fetchone():
            try:
                conn.execute("INSERT OR IGNORE INTO graduate_totals (college, year, total, updated_at, updated_by) "
                             "SELECT ?, year, total, updated_at, updated_by FROM graduate_counts", (DEFAULT_COLLEGE,))
                conn.execute("DROP TABLE IF EXISTS graduate_counts")
            except sqlite3.OperationalError as e:
                # Another process (e.g. Flask's debug reloader) moved them first
                if "no such table" not in str(e):
                    raise
        # Event postings shown on the landing page
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                event_date TEXT NOT NULL,
                event_time TEXT,
                location TEXT,
                description TEXT NOT NULL,
                link TEXT,
                image TEXT,
                published INTEGER NOT NULL DEFAULT 1,
                created_at TEXT,
                created_by TEXT,
                updated_at TEXT
            )
            """
        )
        # Every email sent from the "Email Alumni" page
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS mail_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sent_at TEXT,
                sent_by TEXT,
                recipient_email TEXT,
                recipient_name TEXT,
                subject TEXT,
                status TEXT,
                error TEXT
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS tracer (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                fname TEXT
            )
            """
        )
        # Events and sent emails made before the other colleges were added belong to CCI
        for table in ("events", "mail_log"):
            if "college" not in {row["name"] for row in conn.execute(f"PRAGMA table_info({table})")}:
                try:
                    conn.execute(f"ALTER TABLE {table} ADD COLUMN college TEXT NOT NULL DEFAULT '{DEFAULT_COLLEGE}'")
                except sqlite3.OperationalError as e:
                    if "duplicate column" not in str(e):
                        raise
        existing = {row["name"] for row in conn.execute("PRAGMA table_info(tracer)")}
        for field in FIELDS + EXTRA_COLUMNS:
            if field not in existing:
                try:
                    conn.execute(f'ALTER TABLE tracer ADD COLUMN "{field}" TEXT')
                except sqlite3.OperationalError as e:
                    # Another process (e.g. Flask's debug reloader) added it first
                    if "duplicate column" not in str(e):
                        raise
        # Responses saved before the survey asked for the college all came from CCI
        conn.execute("UPDATE tracer SET college = ? WHERE college IS NULL OR TRIM(college) = ''",
                     (COLLEGE_NAMES[DEFAULT_COLLEGE],))
    conn.close()


def college_programs(college):
    """The programs of a college as [(short code, full name), ...], in the order they were added."""
    conn = get_connection()
    try:
        return [(row["code"], row["name"]) for row in
                conn.execute("SELECT code, name FROM programs WHERE college = ? ORDER BY id", (college,))]
    finally:
        conn.close()


def survey_colleges():
    """For the survey forms: every college with the full names of its programs."""
    return [{"code": code, "name": name, "programs": [program for _, program in college_programs(code)]}
            for code, name in COLLEGES]


def get_answer(field):
    """Join all values sent for a field; 'Other' becomes 'Other: <what they typed>'."""
    other_text = request.form.get(field + "Other", "").strip()
    values = []
    for value in request.form.getlist(field):
        value = value.strip()
        if value == "Other" and other_text:
            value = f"Other: {other_text}"
        values.append(value)
    return "; ".join(values)


def save_proof_file():
    """Save the uploaded proof of employment and return its stored file name ('' if none)."""
    upload = request.files.get("proofFile")
    if not upload or not upload.filename:
        return ""
    extension = os.path.splitext(upload.filename)[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        return ""
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    # Random prefix so two alumni uploading "appointment.pdf" don't overwrite each other
    filename = f"{uuid.uuid4().hex[:8]}_{secure_filename(upload.filename)}"
    upload.save(os.path.join(UPLOAD_DIR, filename))
    return filename


# ---------------------------------------------------------------------------
# Admin accounts
# ---------------------------------------------------------------------------
ROLE_LABELS = {"admin": "System Administrator", "coordinator": "College Coordinator"}


def get_current_admin():
    """Return the logged-in admin as a dict (with 'initials' for the profile badge), or None."""
    admin_id = session.get("admin_id")
    if not admin_id:
        return None
    conn = get_connection()
    try:
        row = conn.execute("SELECT id, full_name, email, mobile, role, active, college FROM admins WHERE id = ?",
                           (admin_id,)).fetchone()
    finally:
        conn.close()
    # A deactivated account is logged out on its next click
    if row is None or not row["active"]:
        return None
    admin = dict(row)
    admin["is_admin"] = admin["role"] == "admin"
    admin["role_label"] = ROLE_LABELS[admin["role"]]
    # The college whose data this account works with: a coordinator's own college, or the one the
    # administrator picked in the header (see /admin/college)
    college = session.get("college") if admin["is_admin"] else admin["college"]
    admin["college"] = college if college in COLLEGE_NAMES else DEFAULT_COLLEGE
    admin["college_name"] = COLLEGE_NAMES[admin["college"]]
    admin["initials"] = "".join(word[0] for word in admin["full_name"].split()[:2]).upper()
    return admin


@app.context_processor
def inject_current_admin():
    # Makes {{ current_admin }} (used by _nav.html) and {{ now_year }} (used by _footer.html) available in every template
    return {"current_admin": get_current_admin(), "now_year": datetime.now().year, "colleges": COLLEGES}


def admin_count():
    conn = get_connection()
    try:
        return conn.execute("SELECT COUNT(*) FROM admins").fetchone()[0]
    finally:
        conn.close()


def admin_required(view):
    """Send visitors to the login page unless an admin is logged in; the admin is in g.admin.

    g.college is the code of the college whose data the page shows.
    """
    @wraps(view)
    def wrapper(*args, **kwargs):
        g.admin = get_current_admin()
        if g.admin is None:
            session.pop("admin_id", None)
            return redirect(url_for("admin_login", next=request.path))
        g.college = g.admin["college"]
        return view(*args, **kwargs)
    return wrapper


def system_admin_required(view):
    """Like admin_required, but only the system administrator may open the page (coordinators get a 403)."""
    @wraps(view)
    @admin_required
    def wrapper(*args, **kwargs):
        if not g.admin["is_admin"]:
            abort(403)
        return view(*args, **kwargs)
    return wrapper


def password_problem(password, confirm):
    """Return an error message if the new password isn't acceptable, else None."""
    if len(password) < MIN_PASSWORD_LENGTH:
        return f"The password must be at least {MIN_PASSWORD_LENGTH} characters."
    if password != confirm:
        return "The two passwords do not match."
    return None


def mobile_problem(mobile):
    """Return an error message if the mobile number doesn't look valid ('' is allowed), else None."""
    digits = re.sub(r"\D", "", mobile)
    if mobile and (not re.fullmatch(r"[0-9+\-\s()]+", mobile) or not 7 <= len(digits) <= 15):
        return "Please enter a valid mobile number, e.g. 09171234567."
    return None


def mask_email(email):
    """'juan.delacruz@gmail.com' -> 'j***@gmail.com'"""
    name, _, domain = email.partition("@")
    return f"{name[:1]}***@{domain}"


def send_email(to, subject, body):
    """Send a plain-text email. Returns False when email isn't set up (the message is printed instead)."""
    if not MAIL_SERVER:
        print(f"\n----- EMAIL NOT SENT (MAIL_SERVER is not set) -----\nTo: {to}\nSubject: {subject}\n\n{body}\n"
              "----------------------------------------------------\n", flush=True)
        return False
    message = EmailMessage()
    message["From"] = MAIL_FROM
    message["To"] = to
    message["Subject"] = subject
    message.set_content(body)
    with smtplib.SMTP(MAIL_SERVER, MAIL_PORT, timeout=20) as smtp:
        smtp.starttls()
        if MAIL_USERNAME:
            smtp.login(MAIL_USERNAME, MAIL_PASSWORD)
        smtp.send_message(message)
    return True


# ---------------------------------------------------------------------------
# Filters for the admin pages, e.g. /admin?year=2020&program=BSIT
# ---------------------------------------------------------------------------
# The year a respondent graduated: the dropdown answer, or else the "Year Graduated (Others)" text
YEAR_SQL = "COALESCE(NULLIF(TRIM(yeargrad), ''), TRIM(yeargradOther))"
NO_FILTERS = {"college": "", "year": "", "program": "", "program_name": "", "type": "", "active": False,
              "query": ""}


def get_filters(default_type=""):
    """Read the year/program/respondent-type filter from the URL. Unknown values are ignored.

    The college is not in the URL: it always comes from the logged-in account (g.college).

    default_type is used when the URL has no type (the All Responses page shows Alumni first).
    """
    codes = dict(college_programs(g.college))
    year = request.args.get("year", "").strip()
    program = request.args.get("program", "").strip()
    rtype = request.args.get("type", "").strip() or default_type
    year = year if re.fullmatch(r"\d{4}", year) else ""
    program = program if program in codes else ""
    rtype = rtype if rtype in RESPONDENT_TYPES else default_type
    query = {key: value for key, value in (("type", rtype), ("year", year), ("program", program)) if value}
    return {
        "college": g.college,
        "year": year,
        "program": program,
        "program_name": codes.get(program, ""),
        "type": rtype,
        "active": bool(year or program),   # the respondent type is a tab, not a filter to clear
        "query": urlencode(query),         # to keep the same selection in links to the other admin pages
    }


def filter_sql(filters):
    """WHERE clause and parameters that select the tracer rows matching the filters."""
    parts, params = [], []
    if filters.get("college"):
        parts.append("college = ?")
        params.append(COLLEGE_NAMES[filters["college"]])
    if filters["year"]:
        parts.append(f"{YEAR_SQL} = ?")
        params.append(filters["year"])
    if filters["program"]:
        # Older records saved the short code (e.g. "BSIT") instead of the full program name
        parts.append("TRIM(program) IN (?, ?)")
        params += [filters["program"], filters["program_name"]]
    if filters.get("type") == "alumni":
        # responses saved before the type was recorded came from the alumni survey
        parts.append("COALESCE(NULLIF(TRIM(respondentType), ''), 'Alumni') = 'Alumni'")
    elif filters.get("type") == "new-graduate":
        parts.append("respondentType = 'New Graduate'")
    return (" WHERE " + " AND ".join(parts) if parts else ""), params


def type_tabs(filters, path):
    """The Alumni / New Graduates tabs: label, link (same year/program), count and whether it's selected."""
    tabs = []
    for key, (label, _) in RESPONDENT_TYPES.items():
        where, params = filter_sql(dict(filters, type=key))
        conn = get_connection()
        try:
            count = conn.execute(f"SELECT COUNT(*) FROM tracer{where}", params).fetchone()[0]
        finally:
            conn.close()
        query = {k: v for k, v in (("type", key), ("year", filters["year"]), ("program", filters["program"])) if v}
        tabs.append({"key": key, "label": label, "count": count, "url": f"{path}?{urlencode(query)}",
                     "selected": filters["type"] == key})
    return tabs


def filter_options():
    """Choices for the filter bar: every year with responses or an entered graduate total, and the programs."""
    conn = get_connection()
    try:
        years = {row[0] for row in conn.execute(f"SELECT DISTINCT {YEAR_SQL} FROM tracer WHERE college = ?",
                                                (COLLEGE_NAMES[g.college],))}
        years |= {row[0] for row in conn.execute("SELECT year FROM graduate_totals WHERE college = ?", (g.college,))}
    finally:
        conn.close()
    return {
        "years": sorted((year for year in years if year and re.fullmatch(r"\d{4}", year)), reverse=True),
        "programs": college_programs(g.college),
    }


def year_summary(filters=NO_FILTERS):
    """Responses per year graduated, matched with the graduate totals the admin entered.

    Returns (rows, overall, unspecified):
      rows        - one dict per year: year, graduates (or None), responses, rate (% or None), updated_at, updated_by
      overall     - totals across the years that have a graduate count: graduates, responses, rate
      unspecified - number of responses with no usable year graduated
    Graduate totals cover all programs, so no rate is given while a program filter is on.
    """
    where, params = filter_sql(filters)
    conn = get_connection()
    try:
        answers = conn.execute(f"SELECT {YEAR_SQL} AS year FROM tracer{where}", params).fetchall()
        totals = {row["year"]: row for row in conn.execute("SELECT * FROM graduate_totals WHERE college = ?",
                                                           (filters["college"],))}
    finally:
        conn.close()

    responses = Counter()
    unspecified = 0
    for answer in answers:
        year = answer["year"] or ""
        if re.fullmatch(r"\d{4}", year):
            responses[year] += 1
        else:
            unspecified += 1

    years = {filters["year"]} if filters["year"] else set(responses) | set(totals)
    with_rate = not filters["program"]
    rows = []
    for year in sorted(years):
        entered = totals.get(year)
        graduates = entered["total"] if entered else None
        rows.append({
            "year": year,
            "graduates": graduates,
            "responses": responses[year],
            "rate": responses[year] / graduates * 100 if graduates and with_rate else None,
            "updated_at": entered["updated_at"] if entered else None,
            "updated_by": entered["updated_by"] if entered else None,
        })

    counted = [row for row in rows if row["graduates"]]
    overall_graduates = sum(row["graduates"] for row in counted)
    overall_responses = sum(row["responses"] for row in counted)
    overall = {
        "graduates": overall_graduates,
        "responses": overall_responses,
        "rate": overall_responses / overall_graduates * 100 if overall_graduates and with_rate else None,
    }
    return rows, overall, unspecified


# ---------------------------------------------------------------------------
# Mail merge ("Email Alumni" page)
# ---------------------------------------------------------------------------
# Tags that can be used in the subject and message; each is replaced with that alumnus' answer
MERGE_FIELDS = {
    "first_name": "First name",
    "last_name": "Last name",
    "middle_name": "Middle name",
    "full_name": "First and last name",
    "program": "Course and specialization",
    "section": "Section",
    "year_graduated": "Year graduated",
    "email": "Email address",
    "sender_name": "Your name (the admin sending)",
}
PLACEHOLDER_RE = re.compile(r"\{(\w+)\}")
EMAIL_RE = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")
DRAFT_MARKER = "[Write your message here.]"


def office_name(college):
    """'CCI' -> 'CCI Alumni Affairs and Relations', the name the college's emails are sent under."""
    return f"{college} Alumni Affairs and Relations"


def default_subject(college):
    return f"Greetings from the {office_name(college)} Office"


def default_body(college):
    return f"""Dear {{last_name}},

Greetings from the {office_name(college)} Office of Iloilo Science and Technology University!

Thank you for taking part in the ISAT U Alumni Tracer Study as a graduate of {{program}}, batch {{year_graduated}}.

{DRAFT_MARKER}

Warm regards,

{{sender_name}}
{office_name(college)}
Iloilo Science and Technology University
"""


def mail_recipients(filters):
    """Alumni matching the filters who gave a valid email, one per address (their latest response).

    Returns (recipients, skipped) - skipped is the number of matching responses without a usable email.
    """
    where, params = filter_sql(filters)
    conn = get_connection()
    try:
        rows = conn.execute(
            f"SELECT id, fname, lname, mname, email, program, section, {YEAR_SQL} AS year "
            f"FROM tracer{where} ORDER BY id", params,
        ).fetchall()
    finally:
        conn.close()

    by_email, skipped = {}, 0
    for row in rows:
        email = (row["email"] or "").strip()
        if not EMAIL_RE.fullmatch(email):
            skipped += 1
            continue
        if email.lower() in by_email:
            skipped += 1   # the same person answered more than once; keep only the newest response
        by_email[email.lower()] = row
    recipients = sorted(by_email.values(), key=lambda r: ((r["lname"] or "").lower(), (r["fname"] or "").lower()))
    return recipients, skipped


def merge_values(row, sender_name):
    """The values that replace each {tag} for one alumnus."""
    first = (row["fname"] or "").strip()
    last = (row["lname"] or "").strip()
    middle = (row["mname"] or "").strip()
    if middle.upper() in ("N/A", "NA", "NONE"):
        middle = ""
    program = (row["program"] or "").strip()
    program = dict(college_programs(g.college)).get(program, program)
    return {
        "first_name": first,
        "last_name": last,
        "middle_name": middle,
        "full_name": " ".join(part for part in (first, last) if part),
        "program": program_name(program) if program else "",
        "section": (row["section"] or "").strip(),
        "year_graduated": (row["year"] or "").strip(),
        "email": (row["email"] or "").strip(),
        "sender_name": sender_name,
    }


def fill_template(text, values, single_line=False):
    """Replace {tags} with the alumnus' values; unknown tags are left as they are."""
    def replace(match):
        if match.group(1) not in values:
            return match.group(0)
        value = values[match.group(1)]
        return " ".join(value.split()) if single_line else value
    return PLACEHOLDER_RE.sub(replace, text)


def unknown_tags(*texts):
    return sorted({tag for text in texts for tag in PLACEHOLDER_RE.findall(text)} - set(MERGE_FIELDS))


def send_bulk(messages, reply_to, sender):
    """Send many emails over one connection. messages: list of (email, name, subject, body).

    sender is the name shown in From, e.g. "CCI Alumni Affairs and Relations".

    Returns a list of (email, name, error) where error is None when the email was accepted.
    """
    results = []
    try:
        with smtplib.SMTP(MAIL_SERVER, MAIL_PORT, timeout=30) as smtp:
            smtp.starttls()
            if MAIL_USERNAME:
                smtp.login(MAIL_USERNAME, MAIL_PASSWORD)
            for to, name, subject, body in messages:
                message = EmailMessage()
                message["From"] = formataddr((sender, MAIL_FROM))
                message["To"] = formataddr((name, to))
                message["Reply-To"] = reply_to
                message["Subject"] = subject
                message.set_content(body)
                try:
                    smtp.send_message(message)
                    results.append((to, name, None))
                except (smtplib.SMTPException, ValueError) as e:
                    results.append((to, name, str(e)))
    except (smtplib.SMTPException, OSError) as e:
        done = {to for to, _, _ in results}
        results += [(to, name, f"Could not reach the mail server: {e}")
                    for to, name, _, _ in messages if to not in done]
    return results


def cancel_password_change(admin_id):
    conn = get_connection()
    try:
        conn.execute("DELETE FROM password_change_requests WHERE admin_id = ?", (admin_id,))
        conn.commit()
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# Events
# ---------------------------------------------------------------------------
@app.template_filter("long_date")
def long_date(value):
    """'2026-11-15' -> 'Sunday, November 15, 2026'"""
    try:
        return date.fromisoformat(value).strftime("%A, %B %d, %Y").replace(" 0", " ")
    except (TypeError, ValueError):
        return value or ""


@app.template_filter("date_badge")
def date_badge(value):
    """'2026-11-15' -> {'month': 'NOV', 'day': '15'} for the calendar badge on event cards."""
    try:
        d = date.fromisoformat(value)
        return {"month": d.strftime("%b").upper(), "day": str(d.day), "year": str(d.year)}
    except (TypeError, ValueError):
        return {"month": "", "day": "", "year": ""}


def save_event_image(upload):
    """Check that the upload really is a picture, shrink it and save it as WebP. Returns the new file name.

    Raises ValueError with a message for the admin when the file can't be used.
    """
    extension = os.path.splitext(upload.filename)[1].lower()
    if extension not in EVENT_IMAGE_EXTENSIONS:
        raise ValueError("The image must be a PNG, JPG, GIF or WebP file.")
    try:
        image = Image.open(upload.stream)
        image.load()
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError):
        raise ValueError("That file could not be read as a picture. Please choose another image.")
    image = ImageOps.exif_transpose(image)   # phone photos: apply the camera's rotation
    if image.mode not in ("RGB", "RGBA"):
        image = image.convert("RGBA" if "transparency" in image.info or image.mode in ("LA", "PA") else "RGB")
    image.thumbnail((EVENT_IMAGE_MAX_SIZE, EVENT_IMAGE_MAX_SIZE), Image.Resampling.LANCZOS)
    os.makedirs(EVENT_IMAGE_DIR, exist_ok=True)
    filename = f"{uuid.uuid4().hex}.webp"
    image.save(os.path.join(EVENT_IMAGE_DIR, filename), "WEBP", quality=85, method=6)
    return filename


def delete_event_image(filename):
    if filename:
        path = os.path.join(EVENT_IMAGE_DIR, os.path.basename(filename))
        if os.path.exists(path):
            os.remove(path)


def read_event_form():
    """The event form's values, tidied up."""
    form = request.form
    return {
        "title": form.get("title", "").strip(),
        "event_date": form.get("event_date", "").strip(),
        "event_time": form.get("event_time", "").strip(),
        "location": form.get("location", "").strip(),
        "description": form.get("description", "").replace("\r\n", "\n").strip(),
        "link": form.get("link", "").strip(),
        "published": 1 if form.get("published") else 0,
    }


def event_problem(values):
    """Return an error message if the event can't be saved, else None."""
    if not values["title"]:
        return "Please enter a title."
    if len(values["title"]) > 150:
        return "The title is too long (150 characters at most)."
    try:
        date.fromisoformat(values["event_date"])
    except ValueError:
        return "Please choose the date of the event."
    if not values["description"]:
        return "Please write a description."
    if values["link"] and not re.match(r"https?://\S+$", values["link"]):
        return "The link must be a full web address starting with http:// or https://"
    return None


def get_event(event_id, published_only=False, college=None):
    conn = get_connection()
    try:
        # With a college, only that college's own event is returned (coordinators manage their own postings)
        row = conn.execute(
            "SELECT * FROM events WHERE id = ?" + (" AND published = 1" if published_only else "")
            + (" AND college = ?" if college else ""), (event_id,) + ((college,) if college else ())
        ).fetchone()
    finally:
        conn.close()
    return row


def public_events(upcoming=True, limit=None):
    """Published events from today on (soonest first), or past events (most recent first)."""
    today = date.today().isoformat()
    sql = ("SELECT * FROM events WHERE published = 1 AND event_date >= ? ORDER BY event_date, id" if upcoming
           else "SELECT * FROM events WHERE published = 1 AND event_date < ? ORDER BY event_date DESC, id DESC")
    if limit:
        sql += f" LIMIT {int(limit)}"
    conn = get_connection()
    try:
        return conn.execute(sql, (today,)).fetchall()
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

# 1. Landing page for visitors
@app.route("/")
def home():
    # Every college that has listed its programs (Programs page), for the "Degree Programs" section
    college_groups = [(code, name, college_programs(code)) for code, name in COLLEGES]
    college_groups = [group for group in college_groups if group[2]]
    return render_template("home.html", college_groups=college_groups, program_info=PROGRAM_INFO,
                           events=public_events(upcoming=True, limit=3))


# 1b. Public event pages (view only)
@app.route("/events")
def events_page():
    return render_template("events.html", upcoming=public_events(upcoming=True),
                           past=public_events(upcoming=False, limit=12))


@app.route("/events/<int:event_id>")
def event_detail(event_id):
    event = get_event(event_id, published_only=True)
    if event is None:
        abort(404)
    return render_template("event_detail.html", event=event, today=date.today().isoformat())


# Browser tab icon: the Alumni Affairs and Relations logo
@app.route("/favicon.ico")
def favicon():
    return send_from_directory(os.path.join(BASE_DIR, "static", "images"), "aar-logo.png", mimetype="image/png")


# 2. The tracer study form
@app.route("/survey")
def survey():
    # First ask who is answering: a new graduate or an alumnus who graduated years ago
    return render_template("survey_choice.html")


@app.route("/survey/alumni")
def survey_alumni():
    # A coordinator can share /survey/alumni?college=CAS so the college is already chosen
    return render_template("index.html", survey_colleges=survey_colleges(),
                           chosen_college=COLLEGE_NAMES.get(request.args.get("college", "").upper(), ""))


@app.route("/survey/new-graduate")
def survey_new_graduate():
    # Alumni Registration Form for graduating students (from the office's Google Form)
    return render_template("survey_new_graduate.html", survey_colleges=survey_colleges(),
                           chosen_college=COLLEGE_NAMES.get(request.args.get("college", "").upper(), ""))


# 3. Route to receive the HTML form data and save it to the DB
@app.route("/submit-data", methods=["POST"])
def submit_data():
    # Every response belongs to one college; that decides which coordinator sees it
    if request.form.get("college", "").strip() not in COLLEGE_NAMES.values():
        abort(400)
    columns = FIELDS + EXTRA_COLUMNS
    values = [get_answer(field) for field in FIELDS]
    values += [save_proof_file(), datetime.now().strftime("%Y-%m-%d %H:%M:%S")]

    column_list = ", ".join(f'"{column}"' for column in columns)
    placeholders = ", ".join("?" for _ in columns)

    conn = get_connection()
    try:
        conn.execute(f"INSERT INTO tracer ({column_list}) VALUES ({placeholders})", values)
        conn.commit()
    finally:
        conn.close()

    return render_template("thanks.html", name=request.form.get("fname", "").strip() or "alumnus")


# 4. Admin account: first-time setup, login and logout
@app.route("/admin/setup", methods=["GET", "POST"])
def admin_setup():
    # Only available while no admin account exists yet
    if admin_count():
        return redirect(url_for("admin_login"))

    error = None
    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        email = request.form.get("email", "").strip().lower()
        mobile = request.form.get("mobile", "").strip()
        password = request.form.get("password", "")
        if not full_name or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            error = "Please enter your full name and a valid email address."
        else:
            error = mobile_problem(mobile) or password_problem(password, request.form.get("confirm", ""))
        if not error:
            conn = get_connection()
            try:
                cursor = conn.execute(
                    "INSERT INTO admins (full_name, email, mobile, password_hash, created_at, role) "
                    "VALUES (?, ?, ?, ?, ?, 'admin')",
                    (full_name, email, mobile, generate_password_hash(password),
                     datetime.now().isoformat(timespec="seconds")),
                )
                conn.commit()
            finally:
                conn.close()
            session.clear()
            session["admin_id"] = cursor.lastrowid
            flash("Your administrator account has been created. You can now register the college coordinators.",
                  "success")
            return redirect(url_for("admin_summary"))
    return render_template("admin_setup.html", error=error, form=request.form, min_length=MIN_PASSWORD_LENGTH)


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if not admin_count():
        return redirect(url_for("admin_setup"))

    next_url = request.values.get("next", "")
    # Only allow redirects back into this site, e.g. "/admin/records"
    if not next_url.startswith("/") or next_url.startswith("//"):
        next_url = url_for("admin_summary")

    error = None
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        conn = get_connection()
        try:
            row = conn.execute("SELECT id, password_hash, active FROM admins WHERE email = ?", (email,)).fetchone()
        finally:
            conn.close()
        password = request.form.get("password", "")
        # Also accept the password without spaces pasted around it (common when copying a temporary password)
        if not row or not (check_password_hash(row["password_hash"], password)
                           or check_password_hash(row["password_hash"], password.strip())):
            error = "Incorrect email or password."
        elif not row["active"]:
            error = "This account has been deactivated. Please contact the system administrator."
        else:
            session.clear()
            session["admin_id"] = row["id"]
            return redirect(next_url)
    return render_template("admin_login.html", error=error, next_url=next_url, email=request.form.get("email", ""))


@app.route("/admin/logout")
def admin_logout():
    session.pop("admin_id", None)
    return redirect(url_for("home"))


# 5. Admin summary of the responses
@app.route("/admin")
@admin_required
def admin_summary():
    filters = get_filters()
    where, params = filter_sql(filters)
    conn = get_connection()
    try:
        rows = conn.execute(f"SELECT program, employed, submittedAt FROM tracer{where}", params).fetchall()
    finally:
        conn.close()

    # Older records saved the short code (e.g. "BSIT"); count both spellings as the same program
    programs = college_programs(g.college)
    code_to_name = {code: name for code, name in programs}
    by_program = {name: {"name": name, "responses": 0, "employed": 0, "not_employed": 0} for _, name in programs}
    for row in rows:
        program = (row["program"] or "").strip()
        program = code_to_name.get(program, program) or "Not specified"
        stats = by_program.setdefault(program, {"name": program, "responses": 0, "employed": 0, "not_employed": 0})
        stats["responses"] += 1
        if row["employed"] == "Yes":
            stats["employed"] += 1
        elif row["employed"] in ("No", "Never employed"):
            stats["not_employed"] += 1
    by_program = list(by_program.values())
    if filters["program"]:
        by_program = [stats for stats in by_program if stats["name"] == filters["program_name"]]
    for stats in by_program:
        stats["chart_row"] = (program_name(stats["name"]), stats["responses"])

    employment = Counter(row["employed"] or "Not specified" for row in rows)
    employment_order = ["Yes", "No", "Never employed"]
    year_rows, overall, year_unspecified = year_summary(filters)

    return render_template(
        "admin.html",
        filters=filters,
        options=filter_options(),
        total=len(rows),
        employed_total=employment.get("Yes", 0),
        latest=max((row["submittedAt"] for row in rows if row["submittedAt"]), default=None),
        by_program=by_program,
        year_rows=year_rows,
        overall=overall,
        year_unspecified=year_unspecified,
        by_employment=sorted(
            employment.items(),
            key=lambda item: employment_order.index(item[0]) if item[0] in employment_order else len(employment_order),
        ),
    )


# 6. Admin list of every saved response
@app.route("/admin/records")
@admin_required
def display():
    filters = get_filters(default_type="alumni")
    where, params = filter_sql(filters)
    conn = get_connection()
    try:
        tracer = conn.execute(f"SELECT * FROM tracer{where} ORDER BY id", params).fetchall()
    finally:
        conn.close()

    type_label, groups = RESPONDENT_TYPES[filters["type"]]
    return render_template("display.html", tracer=tracer, filters=filters, options=filter_options(),
                           groups=groups, type_label=type_label, tabs=type_tabs(filters, "/admin/records"),
                           export_formats=EXPORT_FORMATS)


def export_cell(value):
    """Make a value safe for a spreadsheet: text starting with = + - @ would otherwise run as a formula."""
    if value[:1] in ("=", "+", "-", "@", "\t", "\r") and not re.fullmatch(r"[+-]?[\d\s().-]+", value):
        return "'" + value
    return value


def build_workbook(columns, records, about, groups):
    """Excel file: section headings, column labels, one row per response, plus an 'About' sheet."""
    navy, gold, white = "0D1B8E", "F8C800", "FFFFFF"
    thin = Side(style="thin", color="DDE0EE")
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"

    # Row 1: section names over their columns; row 2: column labels
    ws.cell(row=1, column=1, value="")
    col = 2
    for section, section_columns in groups:
        ws.cell(row=1, column=col, value=section)
        if len(section_columns) > 1:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(section_columns) - 1)
        col += len(section_columns)
    for c in range(1, len(columns) + 1):
        top = ws.cell(row=1, column=c)
        top.fill = PatternFill("solid", fgColor=navy)
        top.font = Font(bold=True, color=white)
        top.alignment = Alignment(horizontal="center", vertical="center")
        head = ws.cell(row=2, column=c, value=columns[c - 1][1])
        head.fill = PatternFill("solid", fgColor=gold)
        head.font = Font(bold=True, color=navy)
        head.alignment = Alignment(wrap_text=True, vertical="center")
        head.border = Border(bottom=thin)

    widths = [len(label) for _, label in columns]
    for r, record in enumerate(records, start=3):
        for c, (key, _) in enumerate(columns, start=1):
            value = record[key]
            if key == "id" or (key in ("yeargrad", "yeargradOther") and re.fullmatch(r"\d{4}", value)):
                cell = ws.cell(row=r, column=c, value=int(value) if value else None)
            elif key == "examRating" and re.fullmatch(r"\d+(\.\d+)?", value):
                cell = ws.cell(row=r, column=c, value=float(value))
            else:
                cell = ws.cell(row=r, column=c, value=value)
                cell.data_type = "s"   # always text: never treated as a formula
                if key == "proofFile" and value:
                    cell.hyperlink = value
                    cell.font = Font(color="2340B8", underline="single")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[c - 1] = max(widths[c - 1], min(len(value), 60))
    for c, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(width + 2, 50))
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "B3"                        # keep the headings and the ID column in view
    last = f"{get_column_letter(len(columns))}{max(len(records) + 2, 2)}"
    ws.auto_filter.ref = f"A2:{last}"             # filter arrows on every column

    info = wb.create_sheet("About this export")
    for r, (label, value) in enumerate(about, start=1):
        info.cell(row=r, column=1, value=label).font = Font(bold=True, color=navy)
        info.cell(row=r, column=2, value=value).data_type = "s"
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# 6a. Admin: download the responses (same filter as the All Responses page)
@app.route("/admin/records/export")
@admin_required
def export_records():
    filters = get_filters(default_type="alumni")
    type_label, groups = RESPONDENT_TYPES[filters["type"]]
    fmt = request.args.get("format", "xlsx")
    if fmt not in EXPORT_FORMATS:
        fmt = "xlsx"
    where, params = filter_sql(filters)
    conn = get_connection()
    try:
        rows = conn.execute(f"SELECT * FROM tracer{where} ORDER BY id", params).fetchall()
    finally:
        conn.close()

    columns = [("id", "ID")] + [column for _, section_columns in groups for column in section_columns]
    records = []
    for row in rows:
        record = {}
        for key, _ in columns:
            value = row[key] if key in row.keys() else None
            value = "" if value is None else str(value)
            if key == "proofFile" and value:
                value = url_for("uploaded_file", filename=value, _external=True)   # opens after admin login
            record[key] = value
        records.append(record)

    now = datetime.now()
    filter_text = f"{filters['year'] or 'All years'} / {filters['program_name'] or 'All programs'}"
    about = [("Exported", now.strftime("%Y-%m-%d %H:%M")), ("Exported by", g.admin["full_name"]),
             ("College", g.admin["college_name"]), ("Respondent type", type_label), ("Filter", filter_text),
             ("Responses", str(len(records))), ("Source", f"ISAT U {g.college} Alumni Tracer Study")]
    filename = "_".join(["tracer-responses", g.college.lower(), filters["type"]] + [part for part in (filters["year"], filters["program"]) if part]
                        + [now.strftime("%Y-%m-%d")]) + "." + fmt

    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([label for _, label in columns])
        for record in records:
            writer.writerow([export_cell(record[key]) for key, _ in columns])
        data = ("\ufeff" + buffer.getvalue()).encode("utf-8")   # the BOM makes Excel read accents like n with tilde
    elif fmt == "json":
        data = json.dumps({
            "exported_at": now.isoformat(timespec="seconds"),
            "exported_by": g.admin["full_name"],
            "college": g.admin["college_name"],
            "respondent_type": type_label,
            "filter": {"year": filters["year"] or None, "program": filters["program_name"] or None},
            "count": len(records),
            "columns": {key: label for key, label in columns},
            "responses": records,
        }, ensure_ascii=False, indent=2).encode("utf-8")
    else:
        data = build_workbook(columns, records, about, groups)

    return Response(data, mimetype=EXPORT_FORMATS[fmt][1],
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


# 6b. Admin: email the alumni (mail merge)
@app.route("/admin/mail", methods=["GET", "POST"])
@admin_required
def admin_mail():
    filters = get_filters()   # the compose form posts back to the same URL, so the filter stays the same
    recipients, skipped = mail_recipients(filters)
    subject, body = default_subject(g.college), default_body(g.college)
    selected = {row["id"] for row in recipients}
    preview = None

    if request.method == "POST":
        subject = request.form.get("subject", "").strip()
        body = request.form.get("body", "").replace("\r\n", "\n")
        selected = {int(i) for i in request.form.getlist("recipient") if i.isdigit()}
        chosen = [row for row in recipients if row["id"] in selected]
        action = request.form.get("action")
        unknown = unknown_tags(subject, body)

        if not subject or not body.strip():
            flash("Please write a subject and a message.", "error")
        elif not chosen:
            flash("Please select at least one recipient.", "error")
        elif action == "preview":
            target = next((row for row in chosen if str(row["id"]) == request.form.get("preview_id")), chosen[0])
            values = merge_values(target, g.admin["full_name"])
            preview = {
                "id": target["id"],
                "to": formataddr((values["full_name"], values["email"])),
                "subject": fill_template(subject, values, single_line=True),
                "body": fill_template(body, values),
            }
            if unknown:
                flash("Unknown tag(s): " + ", ".join("{" + tag + "}" for tag in unknown)
                      + ". They will appear as written. Check the spelling.", "error")
        elif action == "send":
            if not MAIL_SERVER:
                flash("Email sending is not set up yet, so nothing was sent. Set MAIL_SERVER, MAIL_USERNAME and "
                      "MAIL_PASSWORD before starting the app.", "error")
            elif unknown:
                flash("Nothing was sent. Fix the unknown tag(s) first: "
                      + ", ".join("{" + tag + "}" for tag in unknown), "error")
            elif DRAFT_MARKER in body:
                flash(f"Nothing was sent. Replace \"{DRAFT_MARKER}\" with your message first.", "error")
            else:
                messages = []
                for row in chosen:
                    values = merge_values(row, g.admin["full_name"])
                    messages.append((values["email"], values["full_name"],
                                     fill_template(subject, values, single_line=True), fill_template(body, values)))
                results = send_bulk(messages, reply_to=g.admin["email"], sender=office_name(g.college))
                sent_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                conn = get_connection()
                try:
                    conn.executemany(
                        "INSERT INTO mail_log (sent_at, sent_by, recipient_email, recipient_name, subject, status, error, "
                        "college) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                        [(sent_at, g.admin["full_name"], to, name, message[2], "Sent" if error is None else "Failed", error,
                          g.college)
                         for (to, name, error), message in zip(results, messages)],
                    )
                    conn.commit()
                finally:
                    conn.close()
                failed = [(to, error) for to, _, error in results if error]
                flash(f"Sent {len(results) - len(failed)} of {len(results)} email(s).", "success" if not failed else "info")
                if failed:
                    flash("Not sent: " + "; ".join(f"{to} ({error})" for to, error in failed[:5])
                          + (f" and {len(failed) - 5} more" if len(failed) > 5 else ""), "error")
                return redirect(url_for("admin_mail") + ("?" + filters["query"] if filters["query"] else ""))

    conn = get_connection()
    try:
        log = conn.execute("SELECT * FROM mail_log WHERE college = ? ORDER BY id DESC LIMIT 30",
                           (g.college,)).fetchall()
    finally:
        conn.close()
    return render_template(
        "mail.html",
        filters=filters,
        options=filter_options(),
        recipients=recipients,
        skipped=skipped,
        selected=selected,
        subject=subject,
        body=body,
        preview=preview,
        merge_fields=MERGE_FIELDS,
        mail_ready=bool(MAIL_SERVER),
        log=log,
    )


# 6c. Admin: event postings
@app.route("/admin/events")
@admin_required
def admin_events():
    today = date.today().isoformat()
    conn = get_connection()
    try:
        # Upcoming events first (soonest at the top), then past events (most recent first)
        events = conn.execute(
            "SELECT * FROM events WHERE college = ? ORDER BY event_date < ?, "
            "CASE WHEN event_date >= ? THEN event_date END, event_date DESC, id DESC",
            (g.college, today, today),
        ).fetchall()
    finally:
        conn.close()
    return render_template("admin_events.html", events=events, today=today)


@app.route("/admin/events/new", methods=["GET", "POST"])
@app.route("/admin/events/<int:event_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_event(event_id=None):
    event = get_event(event_id, college=g.college) if event_id else None
    if event_id and event is None:
        abort(404)
    values = dict(event) if event else {"title": "", "event_date": "", "event_time": "", "location": "",
                                         "description": "", "link": "", "published": 1, "image": None}
    error = None

    if request.method == "POST":
        values.update(read_event_form())
        error = event_problem(values)
        new_image = None
        upload = request.files.get("image")
        if not error and upload and upload.filename:
            try:
                new_image = save_event_image(upload)
            except ValueError as e:
                error = str(e)

        if not error:
            old_image = event["image"] if event else None
            image = old_image
            if new_image:
                image = new_image
            elif request.form.get("remove_image"):
                image = None
            now = datetime.now().strftime("%Y-%m-%d %H:%M")
            fields = (values["title"], values["event_date"], values["event_time"], values["location"],
                      values["description"], values["link"], image, values["published"])
            conn = get_connection()
            try:
                if event:
                    conn.execute(
                        "UPDATE events SET title = ?, event_date = ?, event_time = ?, location = ?, description = ?, "
                        "link = ?, image = ?, published = ?, updated_at = ? WHERE id = ?",
                        fields + (now, event["id"]),
                    )
                else:
                    conn.execute(
                        "INSERT INTO events (title, event_date, event_time, location, description, link, image, "
                        "published, created_at, created_by, updated_at, college) "
                        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        fields + (now, g.admin["full_name"], now, g.college),
                    )
                conn.commit()
            finally:
                conn.close()
            if old_image and old_image != image:
                delete_event_image(old_image)   # replaced or removed
            state = "published" if values["published"] else "saved as a draft (not shown to visitors)"
            flash(f"\"{values['title']}\" was {state}.", "success")
            return redirect(url_for("admin_events"))

    return render_template("event_form.html", event=event, values=values, error=error)


@app.route("/admin/events/<int:event_id>/delete", methods=["POST"])
@admin_required
def delete_event(event_id):
    event = get_event(event_id, college=g.college)
    if event is None:
        abort(404)
    conn = get_connection()
    try:
        conn.execute("DELETE FROM events WHERE id = ?", (event_id,))
        conn.commit()
    finally:
        conn.close()
    delete_event_image(event["image"])
    flash(f"\"{event['title']}\" was deleted.", "info")
    return redirect(url_for("admin_events"))


# 7. Admin: total number of graduates per year (used for the response rate)
@app.route("/admin/graduates", methods=["GET", "POST"])
@admin_required
def admin_graduates():
    if request.method == "POST":
        year = request.form.get("year", "").strip()
        if not re.fullmatch(r"\d{4}", year) or not 1950 <= int(year) <= datetime.now().year + 1:
            flash("Please enter a valid 4-digit year graduated.", "error")
            return redirect(url_for("admin_graduates"))

        conn = get_connection()
        try:
            if request.form.get("action") == "delete":
                conn.execute("DELETE FROM graduate_totals WHERE college = ? AND year = ?", (g.college, year))
                conn.commit()
                flash(f"Removed the graduate total for {year}.", "info")
                return redirect(url_for("admin_graduates"))

            total = request.form.get("total", "").strip()
            if not total.isdigit() or int(total) < 1:
                flash("The total number of graduates must be a whole number of at least 1.", "error")
                return redirect(url_for("admin_graduates", year=year))
            conn.execute(
                "INSERT INTO graduate_totals (college, year, total, updated_at, updated_by) VALUES (?, ?, ?, ?, ?) "
                "ON CONFLICT(college, year) DO UPDATE SET total = excluded.total, updated_at = excluded.updated_at, "
                "updated_by = excluded.updated_by",
                (g.college, year, int(total), datetime.now().strftime("%Y-%m-%d %H:%M"), g.admin["full_name"]),
            )
            conn.commit()
        finally:
            conn.close()
        flash(f"Saved: {year} has {int(total):,} graduates.", "success")
        return redirect(url_for("admin_graduates"))

    year_rows, _, _ = year_summary(dict(NO_FILTERS, college=g.college))
    edit_year = request.args.get("year", "").strip()
    edit_total = next((row["graduates"] for row in year_rows if row["year"] == edit_year and row["graduates"]), "")
    return render_template("graduates.html", year_rows=year_rows, edit_year=edit_year, edit_total=edit_total)


# 7b. System administrator: register and manage the college coordinators' accounts
def get_coordinator(conn, coordinator_id):
    """Return the coordinator account with this id, or None (the administrator's own account never matches)."""
    return conn.execute("SELECT * FROM admins WHERE id = ? AND role = 'coordinator'", (coordinator_id,)).fetchone()


def new_temporary_password():
    """10 letters and digits that are easy to read out and retype (no 0/O, 1/l/I, and no symbols)."""
    alphabet = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    return "".join(secrets.choice(alphabet) for _ in range(10))


@app.route("/admin/coordinators", methods=["GET", "POST"])
@system_admin_required
def admin_coordinators():
    conn = get_connection()
    try:
        if request.method == "POST":
            action = request.form.get("action", "save")
            coordinator_id = request.form.get("id", "")
            coordinator = get_coordinator(conn, coordinator_id) if coordinator_id else None
            if coordinator_id and coordinator is None:
                flash("That coordinator account no longer exists.", "error")
                return redirect(url_for("admin_coordinators"))

            if action == "save":
                full_name = request.form.get("full_name", "").strip()
                email = request.form.get("email", "").strip().lower()
                mobile = request.form.get("mobile", "").strip()
                password = request.form.get("password", "").strip()   # a stray space would lock the coordinator out
                college = request.form.get("college", "")
                back = url_for("admin_coordinators", edit=coordinator_id or None, _anchor="coordinator-form")
                if not full_name or not EMAIL_RE.fullmatch(email):
                    error = "Please enter the coordinator's full name and a valid email address."
                elif conn.execute("SELECT 1 FROM admins WHERE email = ? AND id != ?",
                                  (email, coordinator_id or 0)).fetchone():
                    error = f"An account with the email {email} already exists."
                elif college not in COLLEGE_NAMES:
                    error = "Please choose the coordinator's college."
                else:
                    error = mobile_problem(mobile)
                if not error and not coordinator and len(password) < MIN_PASSWORD_LENGTH:
                    error = f"The temporary password must be at least {MIN_PASSWORD_LENGTH} characters."
                if error:
                    flash(error, "error")
                    return redirect(back)
                if coordinator:
                    conn.execute("UPDATE admins SET full_name = ?, email = ?, mobile = ?, college = ? WHERE id = ?",
                                 (full_name, email, mobile, college, coordinator["id"]))
                    flash(f"Saved the changes to {full_name}'s account.", "success")
                else:
                    conn.execute(
                        "INSERT INTO admins (full_name, email, mobile, password_hash, created_at, role, college) "
                        "VALUES (?, ?, ?, ?, ?, 'coordinator', ?)",
                        (full_name, email, mobile, generate_password_hash(password),
                         datetime.now().isoformat(timespec="seconds"), college),
                    )
                    flash(f"Registered {full_name} as the {college} coordinator. They log in with {email} and the "
                          "temporary password below, then change it in Settings.", "success")
                    flash(f"Temporary password: {password}", "info")
            elif coordinator is None:
                abort(400)
            elif action == "toggle":
                conn.execute("UPDATE admins SET active = ? WHERE id = ?", (0 if coordinator["active"] else 1,
                                                                           coordinator["id"]))
                if coordinator["active"]:
                    flash(f"{coordinator['full_name']}'s account is deactivated; they can no longer log in.", "info")
                else:
                    flash(f"{coordinator['full_name']}'s account is active again.", "success")
            elif action == "reset":
                password = new_temporary_password()
                conn.execute("UPDATE admins SET password_hash = ? WHERE id = ?",
                             (generate_password_hash(password), coordinator["id"]))
                conn.execute("DELETE FROM password_change_requests WHERE admin_id = ?", (coordinator["id"],))
                flash(f"{coordinator['full_name']}'s password was reset. They log in with {coordinator['email']} and "
                      "the temporary password below, then change it in Settings.", "success")
                flash(f"Temporary password: {password}", "info")
            elif action == "delete":
                conn.execute("DELETE FROM password_change_requests WHERE admin_id = ?", (coordinator["id"],))
                conn.execute("DELETE FROM admins WHERE id = ?", (coordinator["id"],))
                flash(f"Deleted {coordinator['full_name']}'s account.", "info")
            else:
                abort(400)
            conn.commit()
            return redirect(url_for("admin_coordinators"))

        coordinators = conn.execute(
            "SELECT id, full_name, email, mobile, active, created_at, college FROM admins "
            "WHERE role = 'coordinator' ORDER BY active DESC, college, full_name COLLATE NOCASE"
        ).fetchall()
        editing = get_coordinator(conn, request.args.get("edit", ""))
    finally:
        conn.close()
    return render_template("coordinators.html", coordinators=coordinators, editing=editing,
                           college_names=COLLEGE_NAMES, suggested_password=new_temporary_password(), min_length=MIN_PASSWORD_LENGTH)


# 7c. System administrator: choose which college's data to look at (coordinators always see their own)
@app.route("/admin/college", methods=["POST"])
@system_admin_required
def switch_college():
    college = request.form.get("college", "")
    if college in COLLEGE_NAMES:
        session["college"] = college
    # Back to the same page, without the old college's year/program filter
    back = request.form.get("next", "")
    simple_pages = {url_for(name) for name in ("admin_summary", "display", "admin_graduates", "admin_programs",
                                                "admin_mail", "admin_events", "admin_coordinators")}
    return redirect(back if back in simple_pages else url_for("admin_summary"))


# 7d. The college's degree programs (the choices in the survey and in the Program filter)
@app.route("/admin/programs", methods=["GET", "POST"])
@admin_required
def admin_programs():
    college_name = COLLEGE_NAMES[g.college]
    conn = get_connection()
    try:
        if request.method == "POST":
            # Programs can only be added here: removing one would hide it from the survey and the filters
            code = "".join(request.form.get("code", "").upper().split())
            name = " ".join(request.form.get("name", "").upper().split())
            if not re.fullmatch(r"[A-Z0-9][A-Z0-9\-]{1,14}", code):
                flash("The short code must be 2 to 15 letters, numbers or dashes, e.g. BSED-MATH.", "error")
            elif len(name) < 5 or " " not in name:
                flash("Please enter the full name of the program, e.g. Bachelor of Secondary Education.", "error")
            elif conn.execute("SELECT 1 FROM programs WHERE college = ? AND (code = ? OR name = ?)",
                              (g.college, code, name)).fetchone():
                flash(f"{code} or a program with that name is already in the list.", "error")
            else:
                conn.execute("INSERT INTO programs (college, code, name) VALUES (?, ?, ?)", (g.college, code, name))
                conn.commit()
                flash(f"Added {code} to the programs of the {college_name}.", "success")
            return redirect(url_for("admin_programs"))

        programs = conn.execute(
            "SELECT p.id, p.code, p.name, "
            "(SELECT COUNT(*) FROM tracer t WHERE t.college = ? AND TRIM(t.program) IN (p.code, p.name)) AS responses "
            "FROM programs p WHERE p.college = ? ORDER BY p.id", (college_name, g.college)
        ).fetchall()
    finally:
        conn.close()
    return render_template("programs.html", programs=programs)


# 8. Admin settings: profile details and password
@app.route("/admin/settings")
@admin_required
def admin_settings():
    return render_template("settings.html", admin=g.admin, min_length=MIN_PASSWORD_LENGTH)


@app.route("/admin/settings/profile", methods=["POST"])
@admin_required
def update_profile():
    full_name = request.form.get("full_name", "").strip()
    mobile = request.form.get("mobile", "").strip()
    error = "Please enter your full name." if not full_name else mobile_problem(mobile)
    if error:
        flash(error, "error")
    else:
        conn = get_connection()
        try:
            conn.execute("UPDATE admins SET full_name = ?, mobile = ? WHERE id = ?", (full_name, mobile, g.admin["id"]))
            conn.commit()
        finally:
            conn.close()
        flash("Your profile has been updated.", "success")
    return redirect(url_for("admin_settings"))


@app.route("/admin/settings/password", methods=["POST"])
@admin_required
def start_password_change():
    """Check the current password, then email a code; the new password is saved only after the code is entered."""
    current = request.form.get("current_password", "")
    new = request.form.get("new_password", "")

    conn = get_connection()
    try:
        stored_hash = conn.execute("SELECT password_hash FROM admins WHERE id = ?", (g.admin["id"],)).fetchone()[0]
    finally:
        conn.close()

    if not check_password_hash(stored_hash, current):
        error = "Your current password is incorrect."
    elif current == new:
        error = "The new password must be different from your current password."
    else:
        error = password_problem(new, request.form.get("confirm_password", ""))
    if error:
        flash(error, "error")
        return redirect(url_for("admin_settings") + "#password")

    code = f"{secrets.randbelow(1_000_000):06d}"
    expires_at = datetime.now() + CODE_LIFETIME
    conn = get_connection()
    try:
        conn.execute(
            "INSERT OR REPLACE INTO password_change_requests (admin_id, code_hash, new_password_hash, expires_at, attempts) "
            "VALUES (?, ?, ?, ?, 0)",
            (g.admin["id"], generate_password_hash(code), generate_password_hash(new),
             expires_at.isoformat(timespec="seconds")),
        )
        conn.commit()
    finally:
        conn.close()

    body = (
        f"Hello {g.admin['full_name']},\n\n"
        f"Your verification code for changing your ISAT U Alumni Tracer Study password is:\n\n    {code}\n\n"
        f"The code expires in {int(CODE_LIFETIME.total_seconds() // 60)} minutes. "
        "If you did not try to change your password, ignore this email and your password will stay the same.\n"
    )
    try:
        sent = send_email(g.admin["email"], "Your password change verification code", body)
    except (smtplib.SMTPException, OSError) as e:
        cancel_password_change(g.admin["id"])
        flash(f"We could not send the verification email ({e}). Your password was not changed.", "error")
        return redirect(url_for("admin_settings") + "#password")

    if not sent:
        flash("Email sending is not set up yet, so the code was printed in the terminal where Flask is running.", "info")
    return redirect(url_for("verify_password_change"))


@app.route("/admin/settings/verify", methods=["GET", "POST"])
@admin_required
def verify_password_change():
    conn = get_connection()
    try:
        pending = conn.execute("SELECT * FROM password_change_requests WHERE admin_id = ?", (g.admin["id"],)).fetchone()
    finally:
        conn.close()

    if pending is None:
        flash("There is no password change waiting for verification.", "info")
        return redirect(url_for("admin_settings"))
    if datetime.now() > datetime.fromisoformat(pending["expires_at"]):
        cancel_password_change(g.admin["id"])
        flash("The verification code has expired. Please try changing your password again.", "error")
        return redirect(url_for("admin_settings") + "#password")

    error = None
    if request.method == "POST":
        if request.form.get("action") == "cancel":
            cancel_password_change(g.admin["id"])
            flash("Password change cancelled. Your password was not changed.", "info")
            return redirect(url_for("admin_settings"))

        if check_password_hash(pending["code_hash"], request.form.get("code", "").strip()):
            conn = get_connection()
            try:
                conn.execute("UPDATE admins SET password_hash = ? WHERE id = ?",
                             (pending["new_password_hash"], g.admin["id"]))
                conn.execute("DELETE FROM password_change_requests WHERE admin_id = ?", (g.admin["id"],))
                conn.commit()
            finally:
                conn.close()
            try:
                send_email(g.admin["email"], "Your password was changed",
                           f"Hello {g.admin['full_name']},\n\nThe password for your ISAT U Alumni Tracer Study account was just "
                           "changed. If this wasn't you, contact your system administrator right away.\n")
            except (smtplib.SMTPException, OSError):
                pass  # the password is already changed; a missing notice shouldn't undo that
            flash("Your password has been changed.", "success")
            return redirect(url_for("admin_settings"))

        attempts = pending["attempts"] + 1
        if attempts >= MAX_CODE_ATTEMPTS:
            cancel_password_change(g.admin["id"])
            flash("Too many incorrect codes. Your password was not changed; please start again.", "error")
            return redirect(url_for("admin_settings") + "#password")
        conn = get_connection()
        try:
            conn.execute("UPDATE password_change_requests SET attempts = ? WHERE admin_id = ?", (attempts, g.admin["id"]))
            conn.commit()
        finally:
            conn.close()
        left = MAX_CODE_ATTEMPTS - attempts
        error = f"Incorrect code. You have {left} attempt{'s' if left != 1 else ''} left."

    # For the countdown on the page: seconds until the code stops working
    seconds_left = max(0, int((datetime.fromisoformat(pending["expires_at"]) - datetime.now()).total_seconds()))
    return render_template("verify.html", error=error, masked_email=mask_email(g.admin["email"]),
                           seconds_left=seconds_left)


# 9. Route to open an uploaded proof of employment from the records page
@app.route("/uploads/<path:filename>")
@admin_required
def uploaded_file(filename):
    # Only a proof of employment attached to a response of this account's college
    conn = get_connection()
    try:
        owned = conn.execute("SELECT 1 FROM tracer WHERE proofFile = ? AND college = ?",
                             (filename, COLLEGE_NAMES[g.college])).fetchone()
    finally:
        conn.close()
    if not owned:
        abort(404)
    return send_from_directory(UPLOAD_DIR, filename)


# Friendly error pages (instead of Flask's plain text ones)
@app.errorhandler(403)
def not_allowed(error):
    return render_template("error.html", code=403, heading="Not allowed",
                           message="Only the system administrator can open this page."), 403


@app.errorhandler(404)
def page_not_found(error):
    return render_template("error.html", code=404, heading="Page not found",
                           message="Sorry, we couldn't find that page. It may have been moved or removed."), 404


@app.errorhandler(413)
def file_too_large(error):
    limit = app.config["MAX_CONTENT_LENGTH"] // (1024 * 1024)
    return render_template("error.html", code=413, heading="File too large",
                           message=f"The file you attached is larger than {limit} MB. Please go back and choose a "
                                   "smaller file, or a photo or scan with a lower resolution."), 413


@app.errorhandler(500)
def server_error(error):
    return render_template("error.html", code=500, heading="Something went wrong",
                           message="An unexpected error happened on our side. Please try again in a moment. "
                                   "If it keeps happening, contact the Alumni Affairs and Relations Office."), 500


init_db()

if __name__ == "__main__":
    app.run(debug=True)
